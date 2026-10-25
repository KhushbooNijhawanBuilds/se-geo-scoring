"""Consolidate Judge A / Judge B final scorings: re-score, agreement statistics, scored inventory.

Usage: python -I consolidate.py <research_dir> [--no-rescore]

v2.1 (20 Oct 2026):
- Re-scores every judge JSON with the current calculator and weights (skill/geo-research), and rewrites
  <id>_result.json and <id>_score.md, so results always match the standard in force. --no-rescore skips this.
- One method for every page: the FINAL score is Judge A's (the judge that scored all pages). Judge B is a
  blind second opinion on 13 pages, reported beside it with the two-judge mean as a sensitivity check.
  (v2.0 averaged A and B where both existed and used A alone elsewhere, so double-judged types read ~4 points low.)
- Tolerates pages whose capture metadata is missing, and runs on Python 3.8+ (no statistics.correlation).
- Always writes a CSV; also writes an .xlsx when openpyxl is installed.

Outputs (final-scoring/): agreement.json, agreement.md, SE_GEO_scored_inventory_v2.1.csv, SE_GEO_scored_inventory_v2.1.xlsx
"""
import csv, json, os, sys, statistics as st
from collections import defaultdict

R = os.path.abspath(sys.argv[1])
RESCORE = "--no-rescore" not in sys.argv
F = os.path.join(R, "final-scoring")
SKILL = os.path.join(os.path.dirname(R), "skill", "geo-research")
sys.path.insert(0, os.path.join(SKILL, "scripts"))
import geo_score_calc as calc  # noqa: E402

W = calc.load(os.path.join(SKILL, "assets", "weights.json"))
A = json.load(open(os.path.join(F, "assignments.json"), encoding="utf-8"))["assignments"]
rend = {r["n"]: r for r in json.load(open(os.path.join(R, "corpus", "rendered", "index_rendered.json"), encoding="utf-8"))}
BANDS = W["bands"]
band_of = lambda s: next(name for lo, name in BANDS if s >= lo)


def pearson(x, y):
    mx, my = sum(x) / len(x), sum(y) / len(y)
    num = sum((a - mx) * (b - my) for a, b in zip(x, y))
    den = (sum((a - mx) ** 2 for a in x) * sum((b - my) ** 2 for b in y)) ** 0.5
    return num / den if den else None


def load(judge, aid):
    """Return (result, judge_json, legacy_result). Re-scores from the judge JSON when RESCORE is on."""
    j = os.path.join(F, judge, f"{aid}.json")
    p = os.path.join(F, judge, f"{aid}_result.json")
    if not os.path.exists(j):
        return (json.load(open(p, encoding="utf-8")) if os.path.exists(p) else None), None, None
    jj = json.load(open(j, encoding="utf-8"))
    legacy = calc.compute(jj, W, allow_missing=True, legacy=True)
    if RESCORE:
        res = calc.compute(jj, W, allow_missing=True)
        json.dump(res, open(p, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
        open(os.path.join(F, judge, f"{aid}_score.md"), "w", encoding="utf-8").write(calc.to_md(res) + "\n")
    else:
        res = json.load(open(p, encoding="utf-8"))
    return res, jj, legacy


def wkappa(pairs, k=6):
    """Quadratic-weighted Cohen's kappa for ordinal levels 0..5."""
    if len(pairs) < 2:
        return None
    O = [[0] * k for _ in range(k)]
    for a, b in pairs:
        O[a][b] += 1
    n = len(pairs)
    ra = [sum(O[i]) for i in range(k)]
    cb = [sum(O[i][j] for i in range(k)) for j in range(k)]
    num = den = 0.0
    for i in range(k):
        for j in range(k):
            w = (i - j) ** 2 / (k - 1) ** 2
            num += w * O[i][j]
            den += w * ra[i] * cb[j] / n
    return None if den == 0 else round(1 - num / den, 3)


rows, crit_pairs, crit_pairs_by_type, check_agree = [], defaultdict(list), defaultdict(list), [0, 0]
tag_counts = defaultdict(lambda: {"T": 0, "A": 0})
for a in A:
    ra, ja, la_ = load("judgeA", a["id"])
    rb, jb, lb_ = load("judgeB", a["id"])
    row = {k: a[k] for k in ("id", "n", "inventory_type", "title", "url", "type_key", "stage", "target_prompt")}
    row["A"] = ra["final_score"] if ra else None
    row["B"] = rb["final_score"] if rb else None
    row["A_v2_0"] = la_["final_score"] if la_ else None
    row["final"] = row["A"] if row["A"] is not None else row["B"]
    row["final_v2_0_method"] = None  # what the v2.0 consolidation would have reported (A/B mean where paired)
    if la_:
        row["final_v2_0_method"] = round((la_["final_score"] + lb_["final_score"]) / 2) if lb_ else la_["final_score"]
    if ra and rb:
        row["two_judge_mean"] = round((ra["final_score"] + rb["final_score"]) / 2)
        row["diff"] = abs(ra["final_score"] - rb["final_score"])
        row["band_agree"] = ra["band"] == rb["band"]
        lva = {c["id"]: c["level"] for c in ra["criteria"]}
        lvb = {c["id"]: c["level"] for c in rb["criteria"]}
        for cid in set(lva) & set(lvb):
            crit_pairs[cid].append((lva[cid], lvb[cid]))
            crit_pairs_by_type[a["type_key"]].append((lva[cid], lvb[cid]))
        if ja and jb:
            ca = {c["id"]: {k["id"]: k.get("result") for k in c.get("checks", [])} for c in ja.get("criteria", [])}
            cbb = {c["id"]: {k["id"]: k.get("result") for k in c.get("checks", [])} for c in jb.get("criteria", [])}
            for cid in set(ca) & set(cbb):
                for chk in set(ca[cid]) & set(cbb[cid]):
                    check_agree[1] += 1
                    check_agree[0] += ca[cid][chk] == cbb[cid][chk]
    else:
        row["two_judge_mean"] = None
        row["diff"] = None
        row["band_agree"] = None
    ref = ra or rb
    if ref:
        row["band"] = band_of(row["final"])
        row["meets_threshold"] = row["final"] >= W["ok_threshold"]
        row["gates"] = ", ".join(sorted(set((ra or {}).get("gates_triggered", []) + (rb or {}).get("gates_triggered", []))))
        row["penalties"] = ", ".join(sorted(set((ra or {}).get("penalties", []) + (rb or {}).get("penalties", []))))
        top = sorted(ref["criteria"], key=lambda c: -c["impact_points"])[:3]
        row["top_gaps"] = "; ".join(f"{c['id']} (+{c['impact_points']:.1f})" for c in top)
        row["level_changes_v2_1"] = sum(1 for w in ref.get("warnings", []) if " -> " in w and ("anchor:" in w or "rubric cap:" in w))
    for j in (ja, jb):
        if j:
            for c in j.get("criteria", []):
                for chk, t in (c.get("failed_tags") or {}).items():
                    if t in ("T", "A"):
                        tag_counts[a["id"]][t] += 1
    tc = tag_counts[a["id"]]
    row["template_share_of_failed_checks"] = round(tc["T"] / (tc["T"] + tc["A"]), 2) if (tc["T"] + tc["A"]) else None
    rr = rend.get(a["n"]) if not a["id"].endswith("_doc") else None
    if rr and rr.get("server_to_rendered") is not None:
        row["server_words"] = rr.get("server_words")
        row["rendered_words"] = rr.get("words_visible", 0) + rr.get("words_hidden", 0)
        row["server_to_rendered"] = rr["server_to_rendered"]
        row["crawler_gap"] = "YES" if rr["server_to_rendered"] < W["crawler_gap"]["flag_if_server_to_rendered_below"] else "no"
    elif rr:
        row["crawler_gap"] = "unknown (capture metadata missing)"
    else:
        row["crawler_gap"] = "bare PDF" if a["n"] in (59, 60) and not a["id"].endswith("_doc") else "n/a"
    rows.append(row)

paired = [r for r in rows if r["diff"] is not None]
summary = {
    "standard": W.get("version"),
    "final_score_method": "Judge A for every page (single method); Judge B = blind second opinion on the paired pages",
    "scorings": len(rows), "paired": len(paired),
    "mean_abs_diff": round(st.mean(r["diff"] for r in paired), 1) if paired else None,
    "mean_signed_diff_A_minus_B": round(st.mean(r["A"] - r["B"] for r in paired), 1) if paired else None,
    "median_abs_diff": st.median(r["diff"] for r in paired) if paired else None,
    "within_5": sum(r["diff"] <= 5 for r in paired), "within_10": sum(r["diff"] <= 10 for r in paired),
    "band_agreement": round(sum(r["band_agree"] for r in paired) / len(paired), 2) if paired else None,
    "pearson_r": (round(pearson([r["A"] for r in paired], [r["B"] for r in paired]), 3) if len(paired) > 2 else None),
    "check_level_agreement": round(check_agree[0] / check_agree[1], 3) if check_agree[1] else None,
    "criterion_level_wkappa_overall": wkappa([p for v in crit_pairs.values() for p in v]),
    "wkappa_by_type": {t: wkappa(v) for t, v in sorted(crit_pairs_by_type.items())},
    "wkappa_by_criterion": {c: {"kappa": wkappa(v), "n": len(v)} for c, v in sorted(crit_pairs.items())},
    "adjudicate_diff_gt_10": [r["id"] for r in paired if r["diff"] > 10],
    "pages_meeting_threshold": [r["id"] for r in rows if r.get("meets_threshold")],
}
by_type = defaultdict(list)
for r in rows:
    if r["final"] is not None:
        by_type[r["inventory_type"]].append(r["final"])
summary["by_inventory_type"] = {t: {"n": len(v), "min": min(v), "median": st.median(v), "max": max(v)} for t, v in by_type.items()}
json.dump({"summary": summary, "rows": rows}, open(os.path.join(F, "agreement.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)

# Markdown summary
L = [f"# Final scoring {W.get('version')}: Judge A (final) vs blind Judge B", "", "```",
     json.dumps({k: v for k, v in summary.items() if k not in ("wkappa_by_criterion",)}, indent=1, ensure_ascii=False), "```", "",
     "| id | type | stage | final (A) | B | two-judge mean | v2.0 A | band | gates | crawler gap |", "|---|---|---|---|---|---|---|---|---|---|"]
for r in rows:
    L.append(f"| {r['id']} | {r['type_key']} | {r['stage']} | {r['final']} | {r['B'] if r['B'] is not None else ''} | "
             f"{r['two_judge_mean'] if r['two_judge_mean'] is not None else ''} | {r['A_v2_0']} | {r.get('band')} | {r.get('gates', '')} | {r['crawler_gap']} |")
L += ["", "## Weighted kappa by criterion (n pairs)", "", "| criterion | κw | n |", "|---|---|---|"]
for c, v in summary["wkappa_by_criterion"].items():
    L.append(f"| {c} | {v['kappa']} | {v['n']} |")
open(os.path.join(F, "agreement.md"), "w", encoding="utf-8").write("\n".join(L) + "\n")

cols = ["id", "inventory_type", "title", "url", "type_key", "stage", "target_prompt", "final", "band", "meets_threshold", "A", "B",
        "two_judge_mean", "diff", "A_v2_0", "final_v2_0_method", "level_changes_v2_1", "gates", "penalties", "top_gaps",
        "template_share_of_failed_checks", "crawler_gap", "server_words", "rendered_words", "server_to_rendered"]
heads = ["ID", "Page type (inventory)", "Title", "URL", "Scoring logic", "Stage", "Target prompt (inferred)", "GEO score (final, v2.1)", "Band",
         "Meets 70", "Judge A", "Judge B (blind)", "Two-judge mean", "|A−B|", "Judge A under v2.0", "Reported in v2.0", "Levels changed by v2.1 fixes",
         "Gates", "Penalties", "Top 3 gaps (recoverable pts)", "Template share of failed checks", "Crawler gap", "Server-HTML words", "Rendered words", "Server ÷ rendered"]
with open(os.path.join(F, "SE_GEO_scored_inventory_v2.1.csv"), "w", newline="", encoding="utf-8-sig") as fh:
    w = csv.writer(fh)
    w.writerow(heads)
    for r in rows:
        w.writerow([r.get(c) for c in cols])

# Workbook
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Scored inventory"
    ws.append(heads)
    fills = {"Poor": "F8CBAD", "Average": "FFE699", "Good": "C6EFCE", "Excellent": "8FD19E"}
    for r in rows:
        ws.append([r.get(c) for c in cols])
        b = r.get("band")
        if b in fills:
            ws.cell(ws.max_row, 9).fill = PatternFill("solid", fgColor=fills[b])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="3DCD58")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    widths = [7, 22, 40, 40, 18, 13, 40, 10, 10, 8, 8, 9, 9, 7, 9, 9, 9, 14, 18, 36, 12, 14, 10, 10, 10]
    for i, wd in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    s2 = wb.create_sheet("Summary by type")
    s2.append(["Page type", "n", "Min", "Median", "Max"])
    for t, v in summary["by_inventory_type"].items():
        s2.append([t, v["n"], v["min"], v["median"], v["max"]])
    s3 = wb.create_sheet("Agreement")
    for k in ("standard", "final_score_method", "scorings", "paired", "mean_abs_diff", "mean_signed_diff_A_minus_B", "median_abs_diff", "within_5",
              "within_10", "band_agreement", "pearson_r", "check_level_agreement", "criterion_level_wkappa_overall"):
        s3.append([k, summary[k]])
    s3.append([])
    s3.append(["Weighted kappa by type"])
    for t, k in summary["wkappa_by_type"].items():
        s3.append([t, k])
    s3.append([])
    s3.append(["Criterion", "Weighted kappa", "n"])
    for c, v in summary["wkappa_by_criterion"].items():
        s3.append([c, v["kappa"], v["n"]])
    for s in (s2, s3):
        for c in s[1]:
            c.font = Font(bold=True)
        s.column_dimensions["A"].width = 34
    wb.save(os.path.join(F, "SE_GEO_scored_inventory_v2.1.xlsx"))
except ImportError:
    print("openpyxl missing: workbook skipped (CSV written)")
print(json.dumps({k: v for k, v in summary.items() if k != "wkappa_by_criterion"}, indent=1, ensure_ascii=False))
