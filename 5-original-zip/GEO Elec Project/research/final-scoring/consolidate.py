"""Consolidate Judge A / Judge B final scorings: agreement statistics + scored inventory workbook.
Usage: python -E -P consolidate.py <research_dir>
Outputs: final-scoring/agreement.json, final-scoring/agreement.md, final-scoring/SE_GEO_scored_inventory.xlsx
"""
import json, os, sys, statistics as st
from collections import defaultdict

R = sys.argv[1]
F = os.path.join(R, "final-scoring")
A = json.load(open(os.path.join(F, "assignments.json"), encoding="utf-8"))["assignments"]
rend = {r["n"]: r for r in json.load(open(os.path.join(R, "corpus", "rendered", "index_rendered.json"), encoding="utf-8"))}


def load(judge, aid):
    p = os.path.join(F, judge, f"{aid}_result.json")
    j = os.path.join(F, judge, f"{aid}.json")
    if not os.path.exists(p):
        return None, None
    return json.load(open(p, encoding="utf-8")), (json.load(open(j, encoding="utf-8")) if os.path.exists(j) else None)


def wkappa(pairs, k=6):
    """Quadratic-weighted Cohen's kappa for ordinal levels 0..5."""
    if len(pairs) < 2:
        return None
    O = [[0] * k for _ in range(k)]
    for a, b in pairs:
        O[a][b] += 1
    n = len(pairs)
    ra = [sum(O[i]) for i in range(k)]
    cb = [sum(O[i][j] for i in range(k)) for j in range(k)]
    num = den = 0.0
    for i in range(k):
        for j in range(k):
            w = (i - j) ** 2 / (k - 1) ** 2
            num += w * O[i][j]
            den += w * ra[i] * cb[j] / n
    return None if den == 0 else round(1 - num / den, 3)


rows, crit_pairs, crit_pairs_by_type, check_agree = [], defaultdict(list), defaultdict(list), [0, 0]
tag_counts = defaultdict(lambda: {"T": 0, "A": 0})
for a in A:
    ra, ja = load("judgeA", a["id"])
    rb, jb = load("judgeB", a["id"])
    row = {k: a[k] for k in ("id", "n", "inventory_type", "title", "url", "type_key", "stage", "target_prompt")}
    row["A"] = ra["final_score"] if ra else None
    row["B"] = rb["final_score"] if rb else None
    if ra and rb:
        row["final"] = round((ra["final_score"] + rb["final_score"]) / 2)
        row["diff"] = abs(ra["final_score"] - rb["final_score"])
        row["band_agree"] = ra["band"] == rb["band"]
        la = {c["id"]: c["level"] for c in ra["criteria"]}
        lb = {c["id"]: c["level"] for c in rb["criteria"]}
        for cid in set(la) & set(lb):
            crit_pairs[cid].append((la[cid], lb[cid]))
            crit_pairs_by_type[a["type_key"]].append((la[cid], lb[cid]))
        if ja and jb:
            ca = {c["id"]: {k["id"]: k.get("result") for k in c.get("checks", [])} for c in ja.get("criteria", [])}
            cb = {c["id"]: {k["id"]: k.get("result") for k in c.get("checks", [])} for c in jb.get("criteria", [])}
            for cid in set(ca) & set(cb):
                for chk in set(ca[cid]) & set(cb[cid]):
                    check_agree[1] += 1
                    check_agree[0] += ca[cid][chk] == cb[cid][chk]
    else:
        row["final"] = row["A"] if row["A"] is not None else row["B"]
        row["diff"] = None
        row["band_agree"] = None
    ref = ra or rb
    if ref:
        row["band"] = next(name for lo, name in [[90, "Excellent"], [70, "Good"], [40, "Average"], [0, "Poor"]] if row["final"] >= lo)
        row["gates"] = ", ".join(sorted(set((ra or {}).get("gates_triggered", []) + (rb or {}).get("gates_triggered", []))))
        row["penalties"] = ", ".join(sorted(set((ra or {}).get("penalties", []) + (rb or {}).get("penalties", []))))
        imp = defaultdict(list)
        for r in (ra, rb):
            if r:
                for c in r["criteria"]:
                    imp[c["id"]].append(c["impact_points"])
        top = sorted(((st.mean(v), k) for k, v in imp.items()), reverse=True)[:3]
        row["top_gaps"] = "; ".join(f"{k} (+{v:.1f})" for v, k in top)
    for j in (ja, jb):
        if j:
            for c in j.get("criteria", []):
                for chk, t in (c.get("failed_tags") or {}).items():
                    if t in ("T", "A"):
                        tag_counts[a["id"]][t] += 1
    tc = tag_counts[a["id"]]
    row["template_share_of_failed_checks"] = round(tc["T"] / (tc["T"] + tc["A"]), 2) if (tc["T"] + tc["A"]) else None
    rr = rend.get(a["n"]) if not a["id"].endswith("_doc") else None
    if rr:
        row["server_words"] = rr["server_words"]
        row["rendered_words"] = rr["words_visible"] + rr["words_hidden"]
        row["server_to_rendered"] = rr["server_to_rendered"]
        row["crawler_gap"] = "YES" if rr["server_to_rendered"] < 0.7 else "no"
    else:
        row["crawler_gap"] = "bare PDF" if a["n"] in (59, 60) and not a["id"].endswith("_doc") else "n/a"
    rows.append(row)

paired = [r for r in rows if r["diff"] is not None]
summary = {
    "scorings": len(rows), "paired": len(paired),
    "mean_abs_diff": round(st.mean(r["diff"] for r in paired), 1) if paired else None,
    "median_abs_diff": st.median(r["diff"] for r in paired) if paired else None,
    "within_5": sum(r["diff"] <= 5 for r in paired), "within_10": sum(r["diff"] <= 10 for r in paired),
    "band_agreement": round(sum(r["band_agree"] for r in paired) / len(paired), 2) if paired else None,
    "pearson_r": (round(st.correlation([r["A"] for r in paired], [r["B"] for r in paired]), 3) if len(paired) > 2 else None),
    "check_level_agreement": round(check_agree[0] / check_agree[1], 3) if check_agree[1] else None,
    "criterion_level_wkappa_overall": wkappa([p for v in crit_pairs.values() for p in v]),
    "wkappa_by_type": {t: wkappa(v) for t, v in sorted(crit_pairs_by_type.items())},
    "wkappa_by_criterion": {c: {"kappa": wkappa(v), "n": len(v)} for c, v in sorted(crit_pairs.items())},
    "adjudicate_diff_gt_10": [r["id"] for r in paired if r["diff"] > 10],
}
by_type = defaultdict(list)
for r in rows:
    if r["final"] is not None:
        by_type[r["inventory_type"]].append(r["final"])
summary["by_inventory_type"] = {t: {"n": len(v), "min": min(v), "median": st.median(v), "max": max(v)} for t, v in by_type.items()}
json.dump({"summary": summary, "rows": rows}, open(os.path.join(F, "agreement.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)

# Markdown summary
L = ["# Final scoring: inter-rater agreement (Judge A vs blind Judge B)", "", "```", json.dumps({k: v for k, v in summary.items() if k not in ("wkappa_by_criterion",)}, indent=1, ensure_ascii=False), "```", "",
     "| id | type | stage | A | B | final | band | diff | gates | crawler gap |", "|---|---|---|---|---|---|---|---|---|---|"]
for r in rows:
    L.append(f"| {r['id']} | {r['type_key']} | {r['stage']} | {r['A']} | {r['B']} | {r['final']} | {r.get('band')} | {r['diff']} | {r.get('gates','')} | {r['crawler_gap']} |")
L += ["", "## Weighted kappa by criterion (n pairs)", "", "| criterion | κw | n |", "|---|---|---|"]
for c, v in summary["wkappa_by_criterion"].items():
    L.append(f"| {c} | {v['kappa']} | {v['n']} |")
open(os.path.join(F, "agreement.md"), "w", encoding="utf-8").write("\n".join(L))

# Workbook
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Scored inventory"
    cols = ["id", "inventory_type", "title", "url", "type_key", "stage", "target_prompt", "final", "band", "A", "B", "diff",
            "gates", "penalties", "top_gaps", "template_share_of_failed_checks", "crawler_gap", "server_words", "rendered_words", "server_to_rendered"]
    heads = ["ID", "Page type (inventory)", "Title", "URL", "Scoring logic", "Stage", "Target prompt (inferred)", "GEO score (final)", "Band",
             "Judge A", "Judge B (blind)", "|A−B|", "Gates", "Penalties", "Top 3 gaps (recoverable pts)", "Template share of failed checks",
             "Crawler gap", "Server-HTML words", "Rendered words", "Server ÷ rendered"]
    ws.append(heads)
    fills = {"Poor": "F8CBAD", "Average": "FFE699", "Good": "C6EFCE", "Excellent": "8FD19E"}
    for r in rows:
        ws.append([r.get(c) for c in cols])
        b = r.get("band")
        if b in fills:
            ws.cell(ws.max_row, 9).fill = PatternFill("solid", fgColor=fills[b])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="3DCD58")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    widths = [7, 22, 40, 40, 18, 13, 40, 10, 10, 8, 9, 7, 14, 18, 36, 12, 11, 10, 10, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    s2 = wb.create_sheet("Summary by type")
    s2.append(["Page type", "n", "Min", "Median", "Max"])
    for t, v in summary["by_inventory_type"].items():
        s2.append([t, v["n"], v["min"], v["median"], v["max"]])
    s3 = wb.create_sheet("Agreement")
    for k in ("scorings", "paired", "mean_abs_diff", "median_abs_diff", "within_5", "within_10", "band_agreement", "pearson_r", "check_level_agreement", "criterion_level_wkappa_overall"):
        s3.append([k, summary[k]])
    s3.append([])
    s3.append(["Weighted kappa by type"])
    for t, k in summary["wkappa_by_type"].items():
        s3.append([t, k])
    s3.append([])
    s3.append(["Criterion", "Weighted kappa", "n"])
    for c, v in summary["wkappa_by_criterion"].items():
        s3.append([c, v["kappa"], v["n"]])
    for s in (s2, s3):
        for c in s[1]:
            c.font = Font(bold=True)
        s.column_dimensions["A"].width = 34
    wb.save(os.path.join(F, "SE_GEO_scored_inventory.xlsx"))
except ImportError:
    print("openpyxl missing: workbook skipped")
print(json.dumps({k: v for k, v in summary.items() if k != "wkappa_by_criterion"}, indent=1, ensure_ascii=False))
